"""Styled Excel exports: sheets exist, money is in currency units, totals match, empty range works."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import load_workbook
from pg_client_mixin import PgClientTestCase

from app import models, security


def _bearer_headers(user: models.User) -> dict[str, str]:
    data = {
        "sub": user.username,
        "tenant_id": user.tenant_id,
        "provider_id": getattr(user, "provider_id", None),
        "token_version": user.token_version,
    }
    token = security.create_access_token(data, expires_delta=timedelta(minutes=30))
    return {"Authorization": f"Bearer {token}"}


class TestReportXlsx(PgClientTestCase):
    def setUp(self) -> None:
        super().setUp()
        tenant = models.Tenant(name="Informe Pancho")
        self.session.add(tenant)
        self.session.commit()
        self.session.refresh(tenant)
        self.tenant = tenant
        self.owner = models.User(
            username="rx-owner@test.local",
            hashed_password=security.get_password_hash("secret"),
            full_name="Dueño",
            tenant_id=tenant.id,
            role=models.UserRole.owner,
        )
        self.session.add(self.owner)
        floor = models.Floor(name="Main", tenant_id=tenant.id)
        self.session.add(floor)
        self.session.commit()
        self.session.refresh(floor)
        self.session.refresh(self.owner)
        self.table = models.Table(name="T1", tenant_id=tenant.id, floor_id=floor.id, is_active=True)
        self.session.add(self.table)
        self.product = models.Product(name="Café", price_cents=300000, tenant_id=tenant.id, category="Bebidas")
        self.session.add(self.product)
        self.session.commit()
        self.session.refresh(self.table)
        self.session.refresh(self.product)

    def _paid_order(self, qty: int, method: str = "cash", tip: int = 0) -> None:
        order = models.Order(
            table_id=self.table.id,
            tenant_id=self.tenant.id,
            status=models.OrderStatus.paid,
            payment_method=method,
            tip_amount_cents=tip,
            paid_at=datetime.now(timezone.utc),
        )
        self.session.add(order)
        self.session.commit()
        self.session.refresh(order)
        self.session.add(
            models.OrderItem(
                order_id=order.id,
                product_id=self.product.id,
                product_name="Café",
                quantity=qty,
                price_cents=300000,
                cost_cents=100000,
                status=models.OrderItemStatus.pending,
            )
        )
        self.session.commit()

    def _range(self) -> str:
        today = datetime.now(timezone.utc).date()
        return f"from_date={today - timedelta(days=1)}&to_date={today + timedelta(days=1)}"

    def test_sales_export_has_styled_sheets_and_unit_money(self) -> None:
        self._paid_order(2, "cash", tip=50000)
        self._paid_order(1, "terminal")
        r = self.client.get(f"/reports/export?{self._range()}&format=xlsx&lang=es", headers=_bearer_headers(self.owner))
        self.assertEqual(r.status_code, 200, r.text)
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(
            wb.sheetnames[:4], ["Resumen", "Por día", "Productos", "Categorías"]
        )
        self.assertIn("Métodos de pago", wb.sheetnames)
        self.assertIn("Horas y días", wb.sheetnames)
        ws = wb["Productos"]
        header = [c.value for c in ws[4]]
        self.assertIn("Clase ABC", header)
        self.assertIn("Margen", header)
        # revenue 3 × 3000 = 9000 in currency units (not cents)
        rev_col = header.index("Ingresos") + 1
        self.assertEqual(ws.cell(row=5, column=rev_col).value, 9000.0)
        self.assertEqual(ws.cell(row=5, column=header.index("Clase ABC") + 1).value, "A")
        self.assertEqual(ws.freeze_panes, "A5")
        self.assertTrue(ws.auto_filter.ref)
        self.assertEqual(len(wb["Resumen"]._charts), 2)
        pay = wb["Métodos de pago"]
        methods = {pay.cell(row=i, column=1).value for i in (5, 6)}
        self.assertEqual(methods, {"Efectivo", "Datáfono / tarjeta"})

    def test_sales_export_english_and_empty_range(self) -> None:
        r = self.client.get(f"/reports/export?{self._range()}&format=xlsx&lang=en", headers=_bearer_headers(self.owner))
        self.assertEqual(r.status_code, 200, r.text)
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames[0], "Summary")
        self.assertIn("By day", wb.sheetnames)

    def test_detailed_export_sheets_and_totals(self) -> None:
        self._paid_order(2)
        r = self.client.get(f"/reports/export-detailed?{self._range()}&lang=es", headers=_bearer_headers(self.owner))
        self.assertEqual(r.status_code, 200, r.text)
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(
            wb.sheetnames,
            ["Resumen", "Ingresos detallados", "Gastos detallados", "Ingresos vs gastos", "Por proveedor", "Por artículo"],
        )
        inc = wb["Ingresos detallados"]
        header = [c.value for c in inc[4]]
        self.assertEqual(header[0], "Fecha y hora")
        self.assertEqual(inc.cell(row=5, column=header.index("Ingresos") + 1).value, 6000.0)
        self.assertIsInstance(inc.cell(row=5, column=1).value, datetime)
        # totals row under the single data row
        self.assertEqual(inc.cell(row=6, column=1).value, "TOTAL")
        self.assertEqual(inc.cell(row=6, column=header.index("Ingresos") + 1).value, 6000.0)
