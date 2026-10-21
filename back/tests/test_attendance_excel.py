"""Monthly attendance Excel export (reports), including optional staff_ids filter."""
from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone
from openpyxl import load_workbook
from pg_client_mixin import PgClientTestCase

from app import models, security


def _bearer_headers(user: models.User) -> dict[str, str]:
    data = {
        "sub": user.username,
        "tenant_id": user.tenant_id,
        "provider_id": getattr(user, "provider_id", None),
        "token_version": user.token_version,
    }
    token = security.create_access_token(data, expires_delta=timedelta(minutes=30))
    return {"Authorization": f"Bearer {token}"}


class TestAttendanceExcel(PgClientTestCase):
    def setUp(self) -> None:
        super().setUp()
        tenant = models.Tenant(name="Att Excel Test")
        self.session.add(tenant)
        self.session.commit()
        self.session.refresh(tenant)
        self.tenant_id = tenant.id

        self.staff_a = models.User(
            username="att-staff-a@test.local",
            hashed_password=security.get_password_hash("secret"),
            full_name="Staff Alpha",
            tenant_id=self.tenant_id,
            role=models.UserRole.waiter,
        )
        self.staff_b = models.User(
            username="att-staff-b@test.local",
            hashed_password=security.get_password_hash("secret"),
            full_name="Staff Beta",
            tenant_id=self.tenant_id,
            role=models.UserRole.waiter,
        )
        self.admin = models.User(
            username="att-admin@test.local",
            hashed_password=security.get_password_hash("secret"),
            full_name="Admin A",
            tenant_id=self.tenant_id,
            role=models.UserRole.owner,
        )
        self.session.add_all([self.staff_a, self.staff_b, self.admin])
        self.session.commit()
        self.session.refresh(self.staff_a)
        self.session.refresh(self.staff_b)
        self.session.refresh(self.admin)

        # March 2026 — deterministic month for assertions
        base = datetime(2026, 3, 15, 9, 0, tzinfo=timezone.utc)
        for user in (self.staff_a, self.staff_b):
            ws = models.WorkSession(
                tenant_id=self.tenant_id,
                user_id=user.id,
                started_at=base,
                ended_at=base + timedelta(hours=4),
            )
            self.session.add(ws)
        self.session.commit()

    def _employee_rows(self, wb):
        ws = wb["Por empleado"]
        header = [c.value for c in ws[4]]
        rows = {}
        for r in ws.iter_rows(min_row=5, values_only=True):
            if r[1] and r[1] != "TOTAL EQUIPO":
                rows[r[1]] = dict(zip(header, r))
        return header, rows

    def test_attendance_excel_with_staff_ids_filters_and_classifies_hours(self) -> None:
        headers = _bearer_headers(self.admin)
        r = self.client.get(
            "/reports/attendance-excel",
            params={"year": 2026, "month": 3, "staff_ids": self.staff_a.id},
            headers=headers,
        )
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("spreadsheetml", r.headers.get("content-type", ""))
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(
            wb.sheetnames,
            ["Resumen", "Por empleado", "Calendario", "Sesiones", "Semanas", "Alertas", "Reglas y festivos"],
        )
        header, rows = self._employee_rows(wb)
        self.assertEqual(set(rows), {"Staff Alpha"})
        a = rows["Staff Alpha"]
        # 09:00 UTC = 04:00 Bogotá on Sunday 15-03-2026: 2 h night + 2 h day, both on a Sunday
        self.assertEqual(a["Horas totales"], 4.0)
        night = [k for k in header if k and k.split(chr(10))[0] == "RDFN"][0]
        day = [k for k in header if k and k.split(chr(10))[0] == "RDF"][0]
        self.assertEqual(a[night], 2.0)
        self.assertEqual(a[day], 2.0)

    def test_weekly_overtime_above_legal_limit(self) -> None:
        # Monday–Saturday 08:00–16:00 Bogotá (13:00–21:00 UTC) = 48 h in the week of 2-8 Mar 2026 (limit 44 h)
        for d in range(2, 8):
            start = datetime(2026, 3, d, 13, 0, tzinfo=timezone.utc)
            self.session.add(
                models.WorkSession(
                    tenant_id=self.tenant_id, user_id=self.staff_b.id,
                    started_at=start, ended_at=start + timedelta(hours=8),
                )
            )
        self.session.commit()
        r = self.client.get(
            "/reports/attendance-excel",
            params={"year": 2026, "month": 3, "staff_ids": self.staff_b.id},
            headers=_bearer_headers(self.admin),
        )
        self.assertEqual(r.status_code, 200, r.text)
        header, rows = self._employee_rows(load_workbook(io.BytesIO(r.content)))
        b = rows["Staff Beta"]
        self.assertEqual(b["Total extras"], 4.0)
        self.assertEqual(b["Horas totales"], 52.0)

    def test_colombian_holidays_and_limits(self) -> None:
        from datetime import date

        from app.attendance_colombia import colombian_holidays, weekly_limit_hours

        h = colombian_holidays(2026)
        self.assertEqual(h[date(2026, 4, 3)], "Viernes Santo")
        self.assertEqual(h[date(2026, 1, 12)], "Reyes Magos")  # 6-ene moved to Monday
        self.assertIn(date(2026, 5, 18), h)  # Ascensión
        self.assertEqual(weekly_limit_hours(date(2026, 3, 2)), 44)
        self.assertEqual(weekly_limit_hours(date(2026, 8, 3)), 42)

    def test_open_session_is_alerted_not_counted(self) -> None:
        self.session.add(
            models.WorkSession(
                tenant_id=self.tenant_id, user_id=self.staff_a.id,
                started_at=datetime(2026, 3, 20, 14, 0, tzinfo=timezone.utc), ended_at=None,
            )
        )
        self.session.commit()
        r = self.client.get(
            "/reports/attendance-excel",
            params={"year": 2026, "month": 3, "staff_ids": self.staff_a.id},
            headers=_bearer_headers(self.admin),
        )
        wb = load_workbook(io.BytesIO(r.content))
        alerts = [row[2] for row in wb["Alertas"].iter_rows(min_row=5, values_only=True)]
        self.assertIn("Sesión abierta", alerts)
        _, rows = self._employee_rows(wb)
        self.assertEqual(rows["Staff Alpha"]["Horas totales"], 4.0)

    def test_attendance_excel_invalid_staff_id_400(self) -> None:
        headers = _bearer_headers(self.admin)
        r = self.client.get(
            "/reports/attendance-excel",
            params={"year": 2026, "month": 3, "staff_ids": 999_999_999},
            headers=headers,
        )
        self.assertEqual(r.status_code, 400, r.text)
