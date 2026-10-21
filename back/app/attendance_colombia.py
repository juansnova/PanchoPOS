"""
Monthly attendance workbook for Colombia (GET /reports/attendance-excel).

Turns clock-in/out sessions into the figures a Colombian payroll needs: hours split into
day (06:00–19:00) / night (19:00–06:00), Sunday & festivo hours, and overtime above the weekly
legal limit (Ley 2101 de 2021: 42 h/week from 15-07-2026). Everything is computed in the
restaurant's local time zone. It is an *estimate* to prepare payroll, not legal advice.
"""

from __future__ import annotations

from calendar import monthrange
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill

from . import models
from .export_safety import neutralize_workbook
from .report_xlsx import (
    DARK,
    GREEN,
    GREY,
    INT,
    ORANGE,
    ORANGE_SOFT,
    PALETTE,
    _autowidth,
    _banner,
    _bar_chart,
    _col,
    _data_bars,
    _finish,
    _highlights,
    _kpi,
    _section,
    _table,
)

DEFAULT_TZ = "America/Bogota"
HOURS = "0.00"
NIGHT_START = time(19, 0)  # Ley 2466 de 2025: trabajo nocturno 19:00–06:00
NIGHT_END = time(6, 0)
MAX_DAILY_SESSION_HOURS = 12
MAX_WEEKLY_OVERTIME_HOURS = 12

_WEEKDAYS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
_WEEKDAY_LETTER = ["L", "M", "X", "J", "V", "S", "D"]
_MONTHS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
           "octubre", "noviembre", "diciembre"]
_ROLES = {"owner": "Propietario", "admin": "Administrador", "waiter": "Mesero", "kitchen": "Cocina",
          "courier": "Domiciliario", "cashier": "Caja", "manager": "Gerente"}

# category keys -> (short code, header)
CATS = [
    ("ord_day", "HO", "Ordinarias diurnas"),
    ("ord_night", "RN", "Recargo nocturno (19:00–06:00)"),
    ("ord_df_day", "RDF", "Dominical/festivo diurna"),
    ("ord_df_night", "RDFN", "Dominical/festivo nocturna"),
    ("ext_day", "HED", "Extra diurna"),
    ("ext_night", "HEN", "Extra nocturna"),
    ("ext_df_day", "HEDF", "Extra diurna dom./festivo"),
    ("ext_df_night", "HENF", "Extra nocturna dom./festivo"),
]
EXTRA_KEYS = ["ext_day", "ext_night", "ext_df_day", "ext_df_night"]


# ----------------------------------------------------------------------------- Colombian calendar


def _easter(year: int) -> date:
    a, b, c = year % 19, year // 100, year % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    month = (h + l - 7 * m + 114) // 31
    day = ((h + l - 7 * m + 114) % 31) + 1
    return date(year, month, day)


def _next_monday(d: date) -> date:
    return d if d.weekday() == 0 else d + timedelta(days=7 - d.weekday())


def colombian_holidays(year: int) -> dict[date, str]:
    """Festivos de Colombia (Ley 51 de 1983 «Ley Emiliani»)."""
    out: dict[date, str] = {}
    for m, d, name in [(1, 1, "Año Nuevo"), (5, 1, "Día del Trabajo"), (7, 20, "Grito de Independencia"),
                       (8, 7, "Batalla de Boyacá"), (12, 8, "Inmaculada Concepción"), (12, 25, "Navidad")]:
        out[date(year, m, d)] = name
    for m, d, name in [(1, 6, "Reyes Magos"), (3, 19, "San José"), (6, 29, "San Pedro y San Pablo"),
                       (8, 15, "Asunción de la Virgen"), (10, 12, "Día de la Raza"),
                       (11, 1, "Todos los Santos"), (11, 11, "Independencia de Cartagena")]:
        out[_next_monday(date(year, m, d))] = name
    easter = _easter(year)
    out[easter - timedelta(days=3)] = "Jueves Santo"
    out[easter - timedelta(days=2)] = "Viernes Santo"
    out[easter + timedelta(days=43)] = "Ascensión del Señor"
    out[easter + timedelta(days=64)] = "Corpus Christi"
    out[easter + timedelta(days=71)] = "Sagrado Corazón"
    return out


def weekly_limit_hours(on: date) -> int:
    """Jornada máxima semanal (Ley 2101 de 2021): 48 → 47 → 46 → 44 → 42 h."""
    if on >= date(2026, 7, 15):
        return 42
    if on >= date(2025, 7, 15):
        return 44
    if on >= date(2024, 7, 15):
        return 46
    if on >= date(2023, 7, 15):
        return 47
    return 48


def dominical_surcharge_pct(on: date) -> int:
    """Recargo dominical/festivo (Ley 2466 de 2025): 80 % → 90 % → 100 %."""
    if on >= date(2027, 7, 1):
        return 100
    if on >= date(2026, 7, 1):
        return 90
    if on >= date(2025, 7, 1):
        return 80
    return 75


# ----------------------------------------------------------------------------- time maths


@dataclass
class _Piece:
    start: datetime  # naive local
    day: date
    night: bool
    minutes: float
    session_id: int | None = None
    extra: bool = False


def _tz(name: str | None) -> ZoneInfo:
    try:
        return ZoneInfo((name or DEFAULT_TZ).strip() or DEFAULT_TZ)
    except Exception:
        return ZoneInfo(DEFAULT_TZ)


def _local(dt: datetime, tz: ZoneInfo) -> datetime:
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(tz).replace(tzinfo=None)


def _split_day_night(start: datetime, end: datetime, session_id: int | None) -> list[_Piece]:
    """Cut [start, end) at 06:00 / 19:00 / midnight and tag each part day or night."""
    pieces: list[_Piece] = []
    cur = start
    while cur < end:
        day = cur.date()
        if cur.time() < NIGHT_END:
            nxt, night = datetime.combine(day, NIGHT_END), True
        elif cur.time() < NIGHT_START:
            nxt, night = datetime.combine(day, NIGHT_START), False
        else:
            nxt, night = datetime.combine(day + timedelta(days=1), time(0, 0)), True
        seg_end = min(nxt, end)
        mins = (seg_end - cur).total_seconds() / 60
        if mins > 0:
            pieces.append(_Piece(cur, day, night, mins, session_id))
        cur = seg_end
    return pieces


def _subtract(interval: tuple[datetime, datetime], cuts: list[tuple[datetime, datetime]]):
    """Interval minus (possibly overlapping) cut intervals."""
    parts = [interval]
    for cs, ce in sorted(cuts):
        nxt = []
        for ps, pe in parts:
            if ce <= ps or cs >= pe:
                nxt.append((ps, pe))
                continue
            if cs > ps:
                nxt.append((ps, cs))
            if ce < pe:
                nxt.append((ce, pe))
        parts = nxt
    return parts


def _is_df(day: date, holidays: dict[date, str]) -> bool:
    return day.weekday() == 6 or day in holidays


def _cat(piece: _Piece, holidays: dict[date, str]) -> str:
    df = _is_df(piece.day, holidays)
    prefix = "ext" if piece.extra else "ord"
    return f"{prefix}_{'df_' if df else ''}{'night' if piece.night else 'day'}"


def _mark_weekly_overtime(pieces: list[_Piece]) -> list[_Piece]:
    """Hours beyond the weekly legal limit (chronologically) become overtime; pieces are split if needed."""
    by_week: dict[date, list[_Piece]] = defaultdict(list)
    for p in sorted(pieces, key=lambda x: x.start):
        by_week[p.day - timedelta(days=p.day.weekday())].append(p)
    out: list[_Piece] = []
    for monday, items in by_week.items():
        limit = weekly_limit_hours(monday) * 60
        acc = 0.0
        for p in items:
            room = max(0.0, limit - acc)
            if p.minutes <= room:
                acc += p.minutes
                out.append(p)
            elif room <= 0:
                acc += p.minutes
                p.extra = True
                out.append(p)
            else:
                ordinary = _Piece(p.start, p.day, p.night, room, p.session_id)
                extra = _Piece(p.start + timedelta(minutes=room), p.day, p.night, p.minutes - room, p.session_id, True)
                acc += p.minutes
                out.extend([ordinary, extra])
    return out


def _h(minutes: float) -> float:
    return round(minutes / 60, 2)


# ----------------------------------------------------------------------------- workbook


def build_attendance_workbook(
    session,
    tenant: models.Tenant,
    users: list[models.User],
    year: int,
    month: int,
) -> Workbook:
    from sqlmodel import col, select

    tz = _tz(getattr(tenant, "timezone", None))
    first_day = date(year, month, 1)
    last_day = date(year, month, monthrange(year, month)[1])
    ndays = last_day.day
    holidays = {**colombian_holidays(year - 1), **colombian_holidays(year), **colombian_holidays(year + 1)}

    # whole ISO weeks around the month so weekly overtime is right at the edges
    range_start = first_day - timedelta(days=first_day.weekday())
    range_end = last_day + timedelta(days=6 - last_day.weekday())
    start_utc = datetime.combine(range_start, time.min, tzinfo=tz).astimezone(timezone.utc)
    end_utc = datetime.combine(range_end + timedelta(days=1), time.min, tzinfo=tz).astimezone(timezone.utc)
    user_ids = [u.id for u in users]

    sessions = session.exec(
        select(models.WorkSession)
        .where(models.WorkSession.tenant_id == tenant.id)
        .where(col(models.WorkSession.user_id).in_(user_ids))
        .where(models.WorkSession.started_at >= start_utc)
        .where(models.WorkSession.started_at < end_utc)
        .order_by(models.WorkSession.started_at.asc())
    ).all()
    breaks_by_session: dict[int, list[tuple[datetime, datetime]]] = defaultdict(list)
    if sessions:
        for br in session.exec(
            select(models.WorkSessionBreak).where(col(models.WorkSessionBreak.work_session_id).in_([s.id for s in sessions]))
        ).all():
            if br.started_at:
                end = br.ended_at or datetime.now(timezone.utc)
                breaks_by_session[br.work_session_id].append((_local(br.started_at, tz), _local(end, tz)))
    shifts = session.exec(
        select(models.Shift)
        .where(models.Shift.tenant_id == tenant.id)
        .where(col(models.Shift.user_id).in_(user_ids))
        .where(models.Shift.shift_date >= first_day)
        .where(models.Shift.shift_date <= last_day)
    ).all()
    planned_min: dict[int, float] = defaultdict(float)
    for sh in shifts:
        a = datetime.combine(sh.shift_date, sh.start_time)
        b = datetime.combine(sh.shift_date, sh.end_time)
        if b <= a:
            b += timedelta(days=1)
        planned_min[sh.user_id] += (b - a).total_seconds() / 60

    # ---- per-user pieces
    alerts: list[tuple[str, date | None, str, str]] = []
    per_user: dict[int, dict] = {}
    session_rows: list[dict] = []
    for u in users:
        name = u.full_name or u.username or str(u.id)
        pieces: list[_Piece] = []
        n_sessions = 0
        for ws in (s for s in sessions if s.user_id == u.id):
            ls = _local(ws.started_at, tz)
            in_month = first_day <= ls.date() <= last_day
            if ws.ended_at is None:
                if in_month:
                    alerts.append((name, ls.date(), "Sesión abierta",
                                   f"Entró a las {ls:%H:%M} y no registró salida; no se cuenta en las horas."))
                    session_rows.append({"user": u, "name": name, "start": ls, "end": None, "break_min": 0,
                                         "net": 0.0, "day_m": 0.0, "night_m": 0.0, "df": _is_df(ls.date(), holidays),
                                         "note": "Abierta (sin salida)"})
                    n_sessions += 1
                continue
            le = _local(ws.ended_at, tz)
            cuts = breaks_by_session.get(ws.id, [])
            sess_pieces: list[_Piece] = []
            for a, b in _subtract((ls, le), cuts):
                sess_pieces.extend(_split_day_night(a, b, ws.id))
            pieces.extend(sess_pieces)
            if in_month:
                n_sessions += 1
                br_min = sum(max(0.0, (min(e, le) - max(s_, ls)).total_seconds() / 60) for s_, e in cuts)
                net = sum(p.minutes for p in sess_pieces)
                note = ""
                if net / 60 > MAX_DAILY_SESSION_HOURS:
                    note = f"Más de {MAX_DAILY_SESSION_HOURS} h seguidas: revisar"
                    alerts.append((name, ls.date(), "Jornada muy larga", f"{_h(net):.2f} h en una sola sesión."))
                session_rows.append({
                    "user": u, "name": name, "start": ls, "end": le, "break_min": round(br_min),
                    "net": net, "day_m": sum(p.minutes for p in sess_pieces if not p.night),
                    "night_m": sum(p.minutes for p in sess_pieces if p.night),
                    "df": _is_df(ls.date(), holidays), "note": note,
                })
        pieces = _mark_weekly_overtime(pieces)
        in_month_pieces = [p for p in pieces if first_day <= p.day <= last_day]
        totals: dict[str, float] = defaultdict(float)
        daily: dict[date, float] = defaultdict(float)
        for p in in_month_pieces:
            totals[_cat(p, holidays)] += p.minutes
            daily[p.day] += p.minutes
        # weeks
        weeks: dict[date, dict] = {}
        for p in pieces:
            monday = p.day - timedelta(days=p.day.weekday())
            w = weeks.setdefault(monday, {"min": 0.0, "extra": 0.0, "days": set()})
            w["min"] += p.minutes
            w["days"].add(p.day)
            if p.extra:
                w["extra"] += p.minutes
        for monday, w in weeks.items():
            sunday = monday + timedelta(days=6)
            if sunday < first_day or monday > last_day:
                continue
            if len(w["days"]) == 7:
                alerts.append((name, monday, "Sin descanso semanal",
                               f"Trabajó los 7 días de la semana del {monday:%d/%m} al {sunday:%d/%m}."))
            if w["extra"] / 60 > MAX_WEEKLY_OVERTIME_HOURS:
                alerts.append((name, monday, "Extras sobre el máximo",
                               f"{_h(w['extra']):.2f} h extra en la semana del {monday:%d/%m} (máx. {MAX_WEEKLY_OVERTIME_HOURS} h)."))
        per_user[u.id] = {"user": u, "name": name, "totals": totals, "daily": daily, "weeks": weeks,
                          "sessions": n_sessions, "days": len([d for d, m in daily.items() if m > 0])}

    ordered = [per_user[u.id] for u in users]
    team_min = sum(sum(x["totals"].values()) for x in ordered)
    extra_min = sum(sum(x["totals"].get(k, 0) for k in EXTRA_KEYS) for x in ordered)
    night_min = sum(x["totals"].get(k, 0) for x in ordered for k in ("ord_night", "ord_df_night", "ext_night", "ext_df_night"))
    df_min = sum(x["totals"].get(k, 0) for x in ordered for k in ("ord_df_day", "ord_df_night", "ext_df_day", "ext_df_night"))
    tenant_name = getattr(tenant, "name", "") or ""
    period_txt = f"{_MONTHS[month - 1].capitalize()} {year}"
    now_txt = datetime.now(tz).strftime("%d/%m/%Y %H:%M")
    sub = (f"{tenant_name + '  ·  ' if tenant_name else ''}Periodo: {period_txt} "
           f"({first_day:%d/%m/%Y} – {last_day:%d/%m/%Y})  ·  Generado: {now_txt}  ·  Hora local: {tz.key}")

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Resumen"

    # ---------------------------------------------------------------- Por empleado
    ws_e = wb.create_sheet("Por empleado")
    head = ["No. empleado", "Empleado", "Cargo", "Días trabajados", "Sesiones", "Horas totales"]
    head += [f"{code}\n{label}" for _, code, label in CATS]
    head += ["Total extras", "Horas planeadas", "Real vs plan"]
    fm = [None, None, None, INT, INT, HOURS] + [HOURS] * len(CATS) + [HOURS, HOURS, HOURS]
    _banner(ws_e, f"Asistencia · Por empleado · {period_txt}", sub, len(head))
    erows = []
    for x in ordered:
        u, t = x["user"], x["totals"]
        planned = planned_min.get(u.id)
        total = sum(t.values())
        erows.append(
            [u.employee_number or "", x["name"], _ROLES.get(getattr(u.role, "value", str(u.role)), str(getattr(u.role, "value", u.role))),
             x["days"], x["sessions"], _h(total)]
            + [_h(t.get(k, 0)) for k, _, _ in CATS]
            + [_h(sum(t.get(k, 0) for k in EXTRA_KEYS)), _h(planned) if planned else None,
               _h(total - planned) if planned else None]
        )
    tot_row = ["", "TOTAL EQUIPO", "", sum(r[3] for r in erows), sum(r[4] for r in erows), _h(team_min)]
    tot_row += [round(sum(r[6 + i] for r in erows), 2) for i in range(len(CATS))]
    tot_row += [_h(extra_min), round(sum(r[-2] or 0 for r in erows), 2) or None, None]
    e_first, e_last, nxt = _table(ws_e, 4, head, erows, fm, total=tot_row, center=(0,))
    ws_e.row_dimensions[4].height = 48
    _data_bars(ws_e, 6, e_first, e_last)
    _data_bars(ws_e, 6 + len(CATS) + 1, e_first, e_last, "FCA5A5")
    _autowidth(ws_e)
    ws_e.column_dimensions["B"].width = 28
    for j in range(7, 7 + len(CATS) + 3):
        ws_e.column_dimensions[_col(j)].width = 15
    _finish(ws_e, 4, tab="059669")
    ws_e.freeze_panes = "C5"

    # ---------------------------------------------------------------- Calendario
    ws_c = wb.create_sheet("Calendario")
    day_dates = [date(year, month, d) for d in range(1, ndays + 1)]
    head = ["Empleado"] + [f"{d.day}\n{_WEEKDAY_LETTER[d.weekday()]}" for d in day_dates] + ["Total h"]
    fm = [None] + ["0.0"] * ndays + [HOURS]
    _banner(ws_c, f"Asistencia · Calendario de horas · {period_txt}", sub, len(head))
    crows = [[x["name"]] + [(_h(x["daily"].get(d, 0)) or None) for d in day_dates] + [_h(sum(x["daily"].values()))]
             for x in ordered]
    team_by_day = [round(sum(x["daily"].get(d, 0) for x in ordered) / 60, 2) or None for d in day_dates]
    people_by_day = [sum(1 for x in ordered if x["daily"].get(d, 0) > 0) or None for d in day_dates]
    c_first, c_last, nxt = _table(ws_c, 4, head, crows, fm,
                                  total=["Horas del equipo"] + team_by_day + [_h(team_min)])
    for j, label in enumerate(["Personas trabajando"], start=0):
        r = nxt
        ws_c.cell(row=r, column=1, value=label).font = Font(bold=True, color=GREY)
        for i, v in enumerate(people_by_day):
            c = ws_c.cell(row=r, column=2 + i, value=v)
            c.alignment = Alignment(horizontal="center")
            c.font = Font(color=GREY)
    off_fill = PatternFill("solid", fgColor="FED7AA")
    sun_fill = PatternFill("solid", fgColor="FDBA74")
    for i, d in enumerate(day_dates):
        hdr = ws_c.cell(row=4, column=2 + i)
        if d in holidays:
            hdr.fill = PatternFill("solid", fgColor="B91C1C")
        elif d.weekday() == 6:
            hdr.fill = PatternFill("solid", fgColor="C2410C")
        elif d.weekday() == 5:
            hdr.fill = PatternFill("solid", fgColor="57534E")
        if _is_df(d, holidays):
            for r in range(c_first, c_last + 1):
                cell = ws_c.cell(row=r, column=2 + i)
                if cell.value is None:
                    cell.fill = off_fill if d in holidays else sun_fill
    if c_last >= c_first:
        ws_c.conditional_formatting.add(
            f"B{c_first}:{_col(1 + ndays)}{c_last}",
            ColorScaleRule(start_type="num", start_value=0, start_color="FFF7ED",
                           end_type="num", end_value=12, end_color="F97316"),
        )
    ws_c.column_dimensions["A"].width = 28
    for j in range(2, 2 + ndays):
        ws_c.column_dimensions[_col(j)].width = 5.2
    ws_c.column_dimensions[_col(2 + ndays)].width = 10
    ws_c.row_dimensions[4].height = 34
    _finish(ws_c, 4, tab="F59E0B")
    ws_c.freeze_panes = "B5"
    legend_row = nxt + 2
    ws_c.cell(row=legend_row, column=1, value="Rojo = festivo · Naranja oscuro = domingo · Gris = sábado").font = Font(italic=True, size=10, color=GREY)

    # ---------------------------------------------------------------- Sesiones
    ws_s = wb.create_sheet("Sesiones")
    head = ["No. empleado", "Empleado", "Fecha", "Día", "Entrada", "Salida", "Descansos (min)", "Horas netas",
            "Horas diurnas", "Horas nocturnas", "Dom./festivo", "Observación"]
    fm = [None, None, "dd/mm/yyyy", None, "hh:mm", "hh:mm", INT, HOURS, HOURS, HOURS, None, None]
    _banner(ws_s, f"Asistencia · Detalle de sesiones · {period_txt}", sub, len(head))
    srows = []
    for r in sorted(session_rows, key=lambda r: (r["name"].lower(), r["start"])):
        d = r["start"].date()
        srows.append([r["user"].employee_number or "", r["name"], d, _WEEKDAYS[d.weekday()],
                      r["start"].time(), r["end"].time() if r["end"] else "—", r["break_min"], _h(r["net"]),
                      _h(r["day_m"]), _h(r["night_m"]),
                      holidays.get(d, "Domingo") if r["df"] else "", r["note"]])
    s_first, s_last, _ = _table(ws_s, 4, head, srows, fm,
                                total=["", "TOTAL", "", "", "", "", sum(r[6] for r in srows),
                                       round(sum(r[7] for r in srows), 2), round(sum(r[8] for r in srows), 2),
                                       round(sum(r[9] for r in srows), 2), "", ""], center=(0, 3, 4, 5))
    _data_bars(ws_s, 8, s_first, s_last)
    _autowidth(ws_s)
    ws_s.column_dimensions["L"].width = 38
    _finish(ws_s, 4, tab="0E7490")

    # ---------------------------------------------------------------- Semanas
    ws_w = wb.create_sheet("Semanas")
    head = ["Empleado", "Semana", "Días trabajados", "Horas trabajadas", "Límite legal (h)", "Horas extra", "Estado"]
    fm = [None, None, INT, HOURS, INT, HOURS, None]
    _banner(ws_w, f"Asistencia · Control semanal · {period_txt}", sub, len(head))
    wrows = []
    for x in ordered:
        for monday in sorted(x["weeks"]):
            sunday = monday + timedelta(days=6)
            if sunday < first_day or monday > last_day:
                continue
            w = x["weeks"][monday]
            limit = weekly_limit_hours(monday)
            state = "Dentro del límite"
            if len(w["days"]) == 7:
                state = "Sin descanso semanal"
            elif w["extra"] / 60 > MAX_WEEKLY_OVERTIME_HOURS:
                state = "Extras sobre el máximo legal"
            elif w["extra"] > 0:
                state = "Con horas extra"
            partial = " (parcial)" if monday < first_day or sunday > last_day else ""
            wrows.append([x["name"], f"{monday:%d/%m} – {sunday:%d/%m}{partial}", len(w["days"]),
                          _h(w["min"]), limit, _h(w["extra"]), state])
    w_first, w_last, _ = _table(ws_w, 4, head, wrows, fm, center=(1,))
    _data_bars(ws_w, 6, w_first, w_last, "FCA5A5")
    _autowidth(ws_w)
    _finish(ws_w, 4, tab="DC2626")

    # ---------------------------------------------------------------- Alertas
    ws_a = wb.create_sheet("Alertas")
    head = ["Empleado", "Fecha", "Alerta", "Detalle"]
    _banner(ws_a, f"Asistencia · Alertas para revisar · {period_txt}", sub, len(head))
    arows = [[n, d, t, det] for n, d, t, det in sorted(alerts, key=lambda a: (a[0].lower(), a[1] or first_day))]
    if not arows:
        arows = [["—", None, "Sin alertas", "No se encontraron sesiones abiertas, jornadas muy largas ni semanas sin descanso."]]
    _table(ws_a, 4, head, arows, [None, "dd/mm/yyyy", None, None], center=(1,))
    _autowidth(ws_a)
    ws_a.column_dimensions["D"].width = 80
    _finish(ws_a, 4, tab="B91C1C")

    # ---------------------------------------------------------------- Reglas
    ws_r = wb.create_sheet("Reglas y festivos")
    _banner(ws_r, "Cómo se calcularon las horas", sub, 2)
    ws_r.column_dimensions["A"].width = 34
    ws_r.column_dimensions["B"].width = 110
    month_holidays = [(d, n) for d, n in sorted(holidays.items()) if first_day <= d <= last_day]
    limit_now = weekly_limit_hours(last_day)
    rules = [
        ("Zona horaria", f"Todas las horas se calculan con la hora local del restaurante ({tz.key})."),
        ("Horas netas", "Tiempo entre la entrada y la salida menos los descansos registrados. Las sesiones sin salida no se cuentan y salen en «Alertas»."),
        ("Diurna / nocturna", "Diurna: 06:00 a 19:00. Nocturna: 19:00 a 06:00 (recargo nocturno del 35 %). Una sesión que cruza el límite se parte automáticamente."),
        ("Dominical y festivo", f"Se marcan las horas trabajadas en domingo o festivo (recargo vigente en este periodo: {dominical_surcharge_pct(last_day)} %)."),
        ("Jornada semanal máxima", f"{limit_now} horas por semana (lunes a domingo) en este periodo. Lo que pase de ese límite se cuenta como hora extra, en orden cronológico."),
        ("Horas extra", "HED extra diurna (+25 %), HEN extra nocturna (+75 %); en domingo o festivo se suman los recargos correspondientes. Máximo legal: 12 horas extra por semana."),
        ("Semanas parciales", "Las semanas que empiezan o terminan fuera del mes se calculan completas para que el límite semanal sea correcto; solo las horas del mes entran a los totales."),
        ("Horas planeadas", "Suma de los turnos del plan de trabajo del mes. «Real vs plan» = horas trabajadas − horas planeadas."),
        ("Códigos", "HO ordinaria diurna · RN recargo nocturno · RDF/RDFN dominical-festivo diurna/nocturna · HED/HEN extra diurna/nocturna · HEDF/HENF extra en dominical-festivo."),
        ("Importante", "Es una estimación para preparar la nómina. Los porcentajes pueden cambiar por ley o por tu contrato; valida con tu contador antes de pagar."),
    ]
    r = 4
    for k, v in rules:
        a = ws_r.cell(row=r, column=1, value=k)
        a.font = Font(bold=True, color=ORANGE)
        a.alignment = Alignment(vertical="top", indent=1)
        b = ws_r.cell(row=r, column=2, value=v)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        ws_r.row_dimensions[r].height = 34
        r += 1
    r += 1
    _section(ws_r, r, f"Festivos de {_MONTHS[month - 1]} {year}", 2)
    r += 1
    if month_holidays:
        for d, n in month_holidays:
            ws_r.cell(row=r, column=1, value=f"{d:%d/%m/%Y} ({_WEEKDAYS[d.weekday()]})").font = Font(bold=True)
            ws_r.cell(row=r, column=2, value=n)
            r += 1
    else:
        ws_r.cell(row=r, column=1, value="Sin festivos en este mes")
    _finish(ws_r, None, tab="78716C")

    # ---------------------------------------------------------------- Resumen
    ws = ws_sum
    _banner(ws, f"Asistencia mensual · {period_txt}", sub, 8)
    for j in range(1, 9):
        ws.column_dimensions[_col(j)].width = 17
    open_sessions = sum(1 for a in alerts if a[2] == "Sesión abierta")
    cards = [
        ("Empleados", len(ordered), INT, "2563EB"),
        ("Sesiones", sum(x["sessions"] for x in ordered), INT, "0E7490"),
        ("Horas trabajadas", _h(team_min), HOURS, ORANGE),
        ("Horas extra", _h(extra_min), HOURS, "B91C1C" if extra_min else GREEN),
        ("Horas nocturnas", _h(night_min), HOURS, "7C3AED"),
        ("Dom. y festivos", _h(df_min), HOURS, "DB2777"),
        ("Festivos del mes", len(month_holidays), INT, "C2410C"),
        ("Alertas", len(alerts), INT, "B91C1C" if alerts else GREEN),
    ]
    for i, (lab, val, fmt, color) in enumerate(cards):
        _kpi(ws, 4 + (i // 4) * 3, 1 + (i % 4) * 2, lab, val, fmt, color)
    hl: list[tuple[str, str, str]] = []
    if ordered and team_min > 0:
        top = max(ordered, key=lambda x: sum(x["totals"].values()))
        hl.append(("Más horas trabajadas", top["name"], f"{_h(sum(top['totals'].values())):.2f} h en {top['days']} días"))
        top_e = max(ordered, key=lambda x: sum(x["totals"].get(k, 0) for k in EXTRA_KEYS))
        e_m = sum(top_e["totals"].get(k, 0) for k in EXTRA_KEYS)
        hl.append(("Más horas extra", top_e["name"] if e_m else "—", f"{_h(e_m):.2f} h extra" if e_m else "Nadie pasó el límite semanal"))
        top_n = max(ordered, key=lambda x: sum(x["totals"].get(k, 0) for k in ("ord_night", "ord_df_night", "ext_night", "ext_df_night")))
        n_m = sum(top_n["totals"].get(k, 0) for k in ("ord_night", "ord_df_night", "ext_night", "ext_df_night"))
        hl.append(("Más horas nocturnas", top_n["name"] if n_m else "—", f"{_h(n_m):.2f} h entre 19:00 y 06:00" if n_m else "Sin trabajo nocturno"))
        busiest = max(day_dates, key=lambda d: sum(x["daily"].get(d, 0) for x in ordered))
        b_min = sum(x["daily"].get(busiest, 0) for x in ordered)
        hl.append(("Día con más horas del equipo", f"{busiest:%d/%m/%Y} ({_WEEKDAYS[busiest.weekday()]})",
                   f"{_h(b_min):.2f} h · {sum(1 for x in ordered if x['daily'].get(busiest, 0) > 0)} personas"))
        hl.append(("Plan vs real", "Con plan de turnos" if planned_min else "Sin plan de turnos",
                   (f"{_h(team_min - sum(planned_min.values())):+.2f} h frente a lo planeado" if planned_min
                    else "Crea turnos en el plan de trabajo para comparar")))
    if not hl:
        hl.append(("—", "—", "Sin horas registradas en este periodo"))
    _section(ws, 10, "Lo más destacado", 8)
    nr = _highlights(ws, 11, {"indicator": "Indicador", "result": "Resultado", "detail": "Detalle"}, hl) + 1
    if erows:
        ch = _bar_chart(ws_e, "Horas trabajadas por empleado", 2, [6], 4, e_first, e_last, w=26, h=9, horizontal=True)
        ws.add_chart(ch, f"A{nr}")
    _finish(ws, None, landscape=False)

    neutralize_workbook(wb)
    return wb
